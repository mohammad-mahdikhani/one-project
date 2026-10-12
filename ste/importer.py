"""Read supplier price lists from Excel, CSV, or the Aquajoy PDF layout."""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .matcher import SheetRow
from .normalize import fold_persian, normalize_code, parse_money

HEADER_ALIASES = {
    "name": {
        "name", "title", "product", "شرح", "شرح کالا", "شرح كالا", "نام",
        "نام کالا", "عنوان", "محصول", "توضیحات",
    },
    "sku": {"sku", "code", "کد", "کد کالا", "کد محصول", "شناسه", "id"},
    "price": {
        "price", "regular_price", "قیمت", "قیمت فروش", "مصرف کننده",
        "مصرف‌کننده", "قیمت مصرف کننده", "قیمت مصرف‌کننده", "پخش",
    },
    "sale_price": {"sale", "sale_price", "قیمت ویژه", "قیمت همکار", "همکار"},
    "stock": {"stock", "qty", "موجودی", "تعداد"},
}

SKIP_NAME = re.compile(r"لیست قیمت|صفحه|page|ردیف", re.I)
CODE_IN_LINE = re.compile(r"(0\d{7}|99\d{6}|02\d{6}|\b\d{8}\b)")
PRICE_IN_LINE = re.compile(r"(\d{1,3}(?:[٬,]\d{3}){1,3})(?!\d)")
PRICE_BARE = re.compile(r"(?<![A-Za-z\-])(\d{6,12})(?!\d)")
SKU_GLUED = re.compile(r"(0\d{7}|99\d{6}|02\d{6})")


def _header_field(cell) -> str | None:
    key = fold_persian(str(cell or "")).replace("‌", " ").strip().lower()
    key = key.replace("كالا", "کالا")
    for field, aliases in HEADER_ALIASES.items():
        if key in aliases:
            return field
    return None


def _row_from_map(index: int, data: dict) -> SheetRow | None:
    name = fold_persian(str(data.get("name") or ""))
    if not name or SKIP_NAME.search(name) or name.isdigit():
        return None
    sku = normalize_code(data.get("sku"))
    try:
        price = parse_money(data.get("price"), allow_empty=True) if "price" in data else None
    except ValueError:
        price = None
    try:
        sale_price = parse_money(data.get("sale_price"), allow_empty=True) if "sale_price" in data else None
    except ValueError:
        sale_price = None
    try:
        stock = parse_money(data.get("stock"), allow_empty=True) if "stock" in data else None
    except ValueError:
        stock = None
    return SheetRow(row=index, name=name, sku=sku, price=price, sale_price=sale_price, stock=stock, raw=data)


def read_xlsx(data: bytes) -> list[SheetRow]:
    book = load_workbook(io.BytesIO(data), data_only=True)
    sheet = book.active
    rows = list(sheet.iter_rows(values_only=True))
    return _rows_from_table(rows)


def read_csv(data: bytes) -> list[SheetRow]:
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:2048]
    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    reader = csv.reader(io.StringIO(text), dialect)
    return _rows_from_table(list(reader))


def _rows_from_table(rows: list) -> list[SheetRow]:
    if not rows:
        raise ValueError("فایل خالی است")
    header_index = 0
    fields: list[str | None] = []
    for index, row in enumerate(rows[:8]):
        mapped = [_header_field(cell) for cell in row]
        if mapped.count("name") and (mapped.count("sku") or mapped.count("price")):
            header_index = index
            fields = mapped
            break
    if not fields:
        # Aquajoy layout: row, code, name, consumer price
        fields = [None, "sku", "name", "price"]
        if rows and _header_field(rows[0][0]) is None and not normalize_code(rows[0][0] if rows[0] else None):
            header_index = 0
            body = rows[1:]
        else:
            header_index = -1
            body = rows
    else:
        body = rows[header_index + 1:]

    items: list[SheetRow] = []
    start = header_index + 2 if header_index >= 0 else 1
    for offset, row in enumerate(body):
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        data: dict = {}
        for field, cell in zip(fields, row):
            if field:
                data[field] = cell
        if "name" not in data or "sku" not in data:
            parsed = parse_loose_row(row)
            if parsed:
                data = {**parsed, **data}
        item = _row_from_map(start + offset, data)
        if item:
            items.append(item)
    if not items:
        raise ValueError("هیچ کالایی در فایل پیدا نشد")
    return items


def parse_loose_row(row) -> dict | None:
    cells = [fold_persian(str(cell)) for cell in row if cell not in (None, "")]
    if len(cells) < 2:
        return None
    sku = None
    price = None
    name_parts = []
    for cell in cells:
        if sku is None and normalize_code(cell) and re.fullmatch(r"\d{6,12}", str(normalize_code(cell))):
            sku = normalize_code(cell)
            continue
        try:
            money = parse_money(cell, allow_empty=True)
        except ValueError:
            money = None
        if money is not None and money >= 1000 and price is None and cell not in {sku}:
            price = money
            continue
        if cell.isdigit() and len(cell) <= 4:
            continue
        name_parts.append(cell)
    name = fold_persian(" ".join(name_parts))
    if not name:
        return None
    return {"sku": sku, "name": name, "price": price}


def extract_sku(text: str) -> str | None:
    compact = fold_persian(text).replace(" ", "")
    match = SKU_GLUED.search(compact) or CODE_IN_LINE.search(fold_persian(text))
    return match.group(1) if match else normalize_code(None)


def extract_price_span(line: str):
    comma = list(PRICE_IN_LINE.finditer(line))
    if comma:
        last = comma[-1]
        try:
            price = parse_money(last.group(1), allow_empty=True)
        except ValueError:
            price = None
        if price and price >= 1000:
            return price, last
    for match in reversed(list(PRICE_BARE.finditer(line))):
        token = match.group(1)
        if SKU_GLUED.fullmatch(token) or CODE_IN_LINE.fullmatch(token):
            continue
        try:
            price = parse_money(token, allow_empty=True)
        except ValueError:
            continue
        if price and price >= 100000:
            return price, match
    return None, None


def _line_product(index: int, line: str) -> SheetRow | None:
    line = fold_persian(line)
    if not line or SKIP_NAME.search(line) or line.startswith("--"):
        return None
    first = line.split()[0] if line.split() else ""
    if _header_field(first) in {"sku", "name", "price"}:
        return None
    price, span = extract_price_span(line)
    if price is None or span is None:
        return None
    name = fold_persian(line[:span.start()])
    sku = extract_sku(name) or extract_sku(line)
    if sku:
        name = fold_persian(re.sub(r"\d{0,3}" + re.escape(sku), " ", name, count=1))
    name = re.sub(r"^\d{1,4}\s+", "", name)
    name = re.sub(r"^(کد|قیمت|مصرف کننده)\s*", "", name)
    return _row_from_map(index, {"sku": sku, "name": name, "price": price})


def read_pdf_text(text: str) -> list[SheetRow]:
    items = read_message_text(text)
    if not items:
        raise ValueError("از این فایل جدولی خوانده نشد")
    return items


def read_message_text(text: str, source: dict | None = None) -> list[SheetRow]:
    items: list[SheetRow] = []
    lines = [fold_persian(line) for line in (text or "").splitlines()]
    pending_name = ""
    for index, line in enumerate(lines, start=1):
        item = _line_product(index, line)
        if item:
            if source:
                item.source = source.get("source", "telegram")
                item.source_url = source.get("url", "")
                item.source_date = source.get("date", "")
                item.raw.update(source)
            items.append(item)
            pending_name = ""
            continue
        if line and pending_name and extract_price_span(line)[0] and not extract_sku(line):
            item = _line_product(index, f"{pending_name} {line}")
            if item:
                if source:
                    item.source = source.get("source", "telegram")
                    item.source_url = source.get("url", "")
                    item.source_date = source.get("date", "")
                    item.raw.update(source)
                items.append(item)
                pending_name = ""
                continue
        if line and not extract_price_span(line)[0]:
            pending_name = line
    return items


def read_pdf(data: bytes) -> list[SheetRow]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ValueError("خواندن PDF روی این سیستم فعال نیست") from exc
    reader = PdfReader(io.BytesIO(data))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    return read_pdf_text(text)


def read_upload(filename: str, data: bytes) -> list[SheetRow]:
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        return read_xlsx(data)
    if name.endswith(".csv") or name.endswith(".tsv"):
        return read_csv(data)
    if name.endswith(".pdf"):
        return read_pdf(data)
    if data[:2] == b"PK":
        return read_xlsx(data)
    if data[:4] == b"%PDF":
        return read_pdf(data)
    return read_csv(data)


def write_template() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "قیمت"
    sheet.append(["ردیف", "کد", "شرح کالا", "مصرف کننده"])
    sheet.append([1, "01103008", "دستگاه تصفیه آب سه مرحله ای", 132000000])
    sheet.append([2, "01105559", "دستگاه تصفیه آب آکواجوی مدل 401 (E) اکونومی", 325000000])
    sheet.append([3, "01305506", "فیلتر کربن نهایی آکواجوی با اتصال LP-QC", 13400000])
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def write_xlsx(rows: list[SheetRow], path: Path | None = None) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "قیمت"
    sheet.append(["ردیف", "کد", "شرح کالا", "مصرف کننده"])
    for index, row in enumerate(rows, start=1):
        sheet.append([index, row.sku or "", row.name, row.price if row.price is not None else "-"])
    buf = io.BytesIO()
    book.save(buf)
    data = buf.getvalue()
    if path:
        path.write_bytes(data)
    return data
